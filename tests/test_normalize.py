from order_workflow.models import SourceType
from order_workflow.steps import normalize


def test_email_with_headers(tmp_path):
    p = tmp_path / "order.eml"
    p.write_text(
        'From: "Acme S.p.A." <a@acme.it>\nTo: b@duferco.example\n'
        "Subject: Ordine 123\nContent-Type: text/plain; charset=utf-8\n\n"
        "corpo del messaggio\n- 25 t tondo\n"
    )
    doc = normalize.run(p)
    assert doc.source_type == SourceType.EMAIL
    assert doc.metadata["subject"] == "Ordine 123"
    assert "Acme" in doc.metadata["from"]
    assert "corpo del messaggio" in doc.text


def test_csv_semicolon_with_preamble(tmp_path):
    p = tmp_path / "order.csv"
    p.write_text(
        "Cliente: Acme S.p.A.\nRif: A-1\n"
        "Codice;Descrizione;Qta;UM;Prezzo\nSKU-1;Cosa;10;t;100,50\n"
    )
    doc = normalize.run(p)
    assert doc.source_type == SourceType.CSV
    assert len(doc.tables) == 1
    assert doc.tables[0].headers == ["Codice", "Descrizione", "Qta", "UM", "Prezzo"]
    assert doc.tables[0].rows == [["SKU-1", "Cosa", "10", "t", "100,50"]]
    assert "Cliente: Acme S.p.A." in doc.text


def test_csv_comma_delimited(tmp_path):
    p = tmp_path / "order.csv"
    p.write_text("Codice,Descrizione,Qty\nSKU-1,Cosa,10\n")
    doc = normalize.run(p)
    assert doc.tables[0].rows == [["SKU-1", "Cosa", "10"]]


def test_xlsx_header_detection(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Letterhead S.r.l."])
    ws.append([])
    ws.append(["Codice", "Descrizione", "Qta"])
    ws.append(["SKU-9", "Roba", 5])
    p = tmp_path / "order.xlsx"
    wb.save(p)

    doc = normalize.run(p)
    assert doc.source_type == SourceType.EXCEL
    assert doc.tables[0].headers[0] == "Codice"
    assert doc.tables[0].rows == [["SKU-9", "Roba", "5"]]
    assert "Letterhead" in doc.text


def test_unsupported_extension(tmp_path):
    p = tmp_path / "order.docx"
    p.write_text("x")
    try:
        normalize.run(p)
        raise AssertionError("expected ValueError")
    except ValueError:
        pass


def test_scanned_pdf_warns_without_ocr(tmp_path):
    # A PDF with no extractable text triggers the OCR fallback path; without
    # OCR dependencies installed the document carries an explicit warning.
    from reportlab.pdfgen import canvas

    p = tmp_path / "scan.pdf"
    c = canvas.Canvas(str(p))
    c.showPage()  # empty page, no text
    c.save()
    doc = normalize.run(p)
    assert doc.text == "" or len(doc.text) < 40
    assert doc.ocr_used or any("OCR" in w for w in doc.warnings)


def test_email_with_csv_attachment(tmp_path):
    import base64

    csv_payload = base64.b64encode(
        b"Codice;Descrizione;Qta;UM;Prezzo\nTND-B450C-12;Tondo;25;t;614,40\n"
    ).decode()
    p = tmp_path / "order.eml"
    p.write_text(
        'From: "Acme S.p.A." <a@acme.it>\nSubject: ordine\nMIME-Version: 1.0\n'
        'Content-Type: multipart/mixed; boundary="BB"\n\n'
        "--BB\nContent-Type: text/plain\n\nordine in allegato\n"
        "--BB\nContent-Type: text/csv\nContent-Disposition: attachment; filename=ordine.csv\n"
        "Content-Transfer-Encoding: base64\n\n" + csv_payload + "\n--BB--\n"
    )
    doc = normalize.run(p)
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[0][0] == "TND-B450C-12"
    assert "ordine in allegato" in doc.text
